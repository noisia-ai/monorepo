#!/usr/bin/env python3
"""
preview.py — render the PPTX this builder makes to PNGs, with real font metrics, and flag
text that does not fit its box.

    python3 preview.py deck.pptx outdir [--fonts <dir>] [--sheet]

It is not a general PowerPoint renderer. It draws the vocabulary noisia_pptx.py uses (slide
background picture, pictures, rounded rectangles, ovals, lines, text boxes with runs and the
slide-number field) and nothing else, which is exactly what makes it reliable for this deck
and cheap enough to run inside the claude.ai sandbox, where there is no Keynote and
LibreOffice substitutes fonts.

Why it exists: a deck checked only in code ships with text cut off. The builder estimates
widths; this measures them with the real Google Sans (or Product Sans, the same design) and
reports every box whose text overflows, so the fix happens before anyone opens the file.
"""
import argparse, glob, io, os, sys
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PX = 6350
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}

FONT_DIRS = [os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "fonts"),
             os.path.expanduser("~/Library/Fonts"), os.path.expanduser("~/.fonts"),
             "/usr/share/fonts", "/Library/Fonts"]


def find_fonts(extra=None):
    dirs = ([extra] if extra else []) + FONT_DIRS
    reg = bold = None
    for d in dirs:
        for f in glob.glob(os.path.join(d, "**", "*.ttf"), recursive=True):
            n = os.path.basename(f).lower().replace(" ", "").replace("-", "").replace("_", "")
            if ("googlesans" in n or "productsans" in n) and "italic" not in n and "code" not in n \
                    and "flex" not in n and "display" not in n:
                if "bold" in n and not bold:
                    bold = f
                elif "regular" in n and not reg:
                    reg = f
    if not reg:
        # Fall back to any sans TTF so the check still runs. A wider font than Google Sans makes
        # the overflow report conservative, never optimistic.
        for d in dirs:
            for f in glob.glob(os.path.join(d, "**", "DejaVuSans*.ttf"), recursive=True) + \
                     glob.glob(os.path.join(d, "**", "LiberationSans*.ttf"), recursive=True):
                n = os.path.basename(f).lower()
                if "bold" in n and "oblique" not in n and "italic" not in n and not bold:
                    bold = f
                elif "bold" not in n and "oblique" not in n and "italic" not in n and "mono" not in n \
                        and "condensed" not in n and not reg:
                    reg = f
        if not reg:
            sys.exit("No encontré Google Sans ni otra fuente sans (TTF). Pasa --fonts <dir>.")
        print("aviso: sin Google Sans, mido con", os.path.basename(reg), "(más ancha: el reporte peca de estricto)")
    return reg, bold or reg


class Fonts:
    def __init__(self, reg, bold):
        self.reg, self.bold, self.cache = reg, bold, {}

    def get(self, px, bold=False):
        k = (round(px * 2) / 2, bold)
        if k not in self.cache:
            self.cache[k] = ImageFont.truetype(self.bold if bold else self.reg, max(1, int(round(k[0]))))
        return self.cache[k]


def hexrgb(h, a=255):
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4)) + (a,)


def shape_px(sh):
    return sh.left / EMU_PX, sh.top / EMU_PX, sh.width / EMU_PX, sh.height / EMU_PX


def blob_image(part, rid):
    return Image.open(io.BytesIO(part.related_part(rid).blob)).convert("RGBA")


def para_runs(p):
    """Runs and fields in order: (text, size_px, bold, rgb)."""
    out = []
    for el in p._p:
        tag = el.tag.split("}")[1]
        if tag not in ("r", "fld"):
            continue
        rpr = el.find("a:rPr", NS)
        t = el.find("a:t", NS)
        txt = t.text if t is not None and t.text else ""
        size = int(rpr.get("sz", "1800")) / 100 * 2 if rpr is not None else 36
        bold = rpr is not None and rpr.get("b") == "1"
        clr = rpr.find(".//a:srgbClr", NS) if rpr is not None else None
        out.append((txt, size, bold, clr.get("val") if clr is not None else "2b2b2b"))
    return out


def layout_paragraph(runs, width, fonts):
    """Greedy word wrap across runs, keeping run boundaries that have no space between them
    (an accent word followed by punctuation). Returns lines: [(segments, max_size)]."""
    import re as _re
    words, pending = [], False
    for txt, size, bold, rgb in runs:
        for tok in _re.findall(r"\S+|\s+", txt):
            if tok.isspace():
                pending = True
                continue
            words.append((tok, size, bold, rgb, pending))
            pending = False
    lines, cur, cur_w, cur_max = [], [], 0, 0
    for w, size, bold, rgb, spaced in words:
        f = fonts.get(size, bold)
        sp = f.getlength(" ") if (cur and spaced) else 0
        ww = f.getlength(w)
        if cur and cur_w + sp + ww > width + 0.5:
            lines.append((cur, cur_max)); cur, cur_w, cur_max = [], 0, 0; sp = 0
        cur.append((w, size, bold, rgb, sp))
        cur_w += sp + ww
        cur_max = max(cur_max, size)
    if cur or not lines:
        lines.append((cur, cur_max or (runs[0][1] if runs else 20)))
    return lines


def draw_text(img, sh, fonts, issues, slide_no):
    x, y, w, h = shape_px(sh)
    tf = sh.text_frame
    d = ImageDraw.Draw(img)
    blocks, total = [], 0
    for p in tf.paragraphs:
        runs = para_runs(p)
        if not runs:
            continue
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        algn = {None: "l", 1: "l", 2: "c", 3: "r"}.get(p.alignment, "l")
        for segs, mx in layout_paragraph(runs, w, fonts):
            lh = mx * 1.2 * ls
            blocks.append((segs, mx, lh, algn))
            total += lh
    if not blocks:
        return
    anchor = tf.vertical_anchor
    oy = y
    if anchor is not None and int(anchor) == 3:      # middle
        oy = y + (h - total) / 2
    elif anchor is not None and int(anchor) == 4:    # bottom
        oy = y + h - total
    if total > h + 2:
        txt = " ".join(r[0] for p in tf.paragraphs for r in para_runs(p))[:60]
        issues.append(f"slide {slide_no}: texto desborda su caja por {total - h:.0f}px · «{txt}…»")
    cy = oy
    for segs, mx, lh, algn in blocks:
        lw = sum(fonts.get(s, b).getlength(t) + sp for t, s, b, c, sp in segs)
        cx = x if algn == "l" else (x + (w - lw) / 2 if algn == "c" else x + w - lw)
        if lw > w + 4:
            issues.append(f"slide {slide_no}: una palabra no cabe a lo ancho («{segs[0][0]}…»)")
        base = cy + (lh - mx * 1.2) / 2 + mx * 0.95
        for t, s, b, c, sp in segs:
            f = fonts.get(s, b)
            cx += sp
            d.text((cx, base), t, font=f, fill=hexrgb(c), anchor="ls")
            cx += f.getlength(t)
        cy += lh


def render_slide(prs, slide, fonts, n, issues):
    W, H = int(prs.slide_width / EMU_PX), int(prs.slide_height / EMU_PX)
    img = Image.new("RGBA", (W, H), (255, 255, 255, 255))
    bg = slide._element.find(".//p:bg//a:blip", NS)
    if bg is not None:
        im = blob_image(slide.part, bg.get("{%s}embed" % NS["r"])).resize((W, H), Image.LANCZOS)
        img.alpha_composite(im)
    for sh in slide.shapes:
        x, y, w, h = shape_px(sh)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA").resize(
                (max(1, int(w)), max(1, int(h))), Image.LANCZOS)
            img.alpha_composite(im, (int(x), int(y))) if x >= 0 and y >= 0 else _paste_clip(img, im, x, y)
        elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            layer = Image.new("RGBA", (W, H), (0, 0, 0, 0))
            d = ImageDraw.Draw(layer)
            fill = None
            if sh.fill.type == 1:
                fill = hexrgb(str(sh.fill.fore_color.rgb))
            line = None
            try:
                if sh.line.fill.type == 1:
                    line = hexrgb(str(sh.line.color.rgb))
            except Exception:
                line = None
            box = [x, y, x + w, y + h]
            name = sh.auto_shape_type
            if "OVAL" in str(name):
                d.ellipse(box, fill=fill, outline=line, width=1)
            elif "ROUNDED" in str(name):
                r = sh.adjustments[0] * min(w, h) if len(sh.adjustments) else 0
                d.rounded_rectangle(box, radius=r, fill=fill, outline=line, width=1)
            else:
                d.rectangle(box, fill=fill, outline=line, width=1)
            img.alpha_composite(layer)
            if sh.has_text_frame and sh.text_frame.text.strip():
                draw_text(img, sh, fonts, issues, n)
        elif sh.shape_type == MSO_SHAPE_TYPE.LINE or sh.__class__.__name__ == "Connector":
            d = ImageDraw.Draw(img)
            try:
                c = hexrgb(str(sh.line.color.rgb))
            except Exception:
                c = hexrgb("eeeeee")
            d.line([sh.begin_x / EMU_PX, sh.begin_y / EMU_PX, sh.end_x / EMU_PX, sh.end_y / EMU_PX], fill=c, width=1)
        elif sh.has_text_frame:
            draw_text(img, sh, fonts, issues, n)
    return img.convert("RGB")


def _paste_clip(img, im, x, y):
    sx, sy = int(max(0, -x)), int(max(0, -y))
    crop = im.crop((sx, sy, im.width, im.height))
    img.alpha_composite(crop, (int(max(0, x)), int(max(0, y))))


def contact_sheet(pngs, path, cols=2, tw=960):
    ims = [Image.open(p) for p in pngs]
    th = int(ims[0].height * tw / ims[0].width)
    rows = (len(ims) + cols - 1) // cols
    sh = Image.new("RGB", (cols * tw + (cols + 1) * 12, rows * th + (rows + 1) * 12), "#3a3a3a")
    for i, im in enumerate(ims):
        r, c = divmod(i, cols)
        sh.paste(im.resize((tw, th), Image.LANCZOS), (12 + c * (tw + 12), 12 + r * (th + 12)))
    sh.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx"); ap.add_argument("outdir")
    ap.add_argument("--fonts"); ap.add_argument("--sheet", action="store_true")
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    fonts = Fonts(*find_fonts(a.fonts))
    prs = Presentation(a.pptx)
    issues, pngs = [], []
    for n, s in enumerate(prs.slides, 1):
        p = os.path.join(a.outdir, f"slide-{n:02d}.png")
        render_slide(prs, s, fonts, n, issues).save(p)
        pngs.append(p)
    if a.sheet:
        for i in range(0, len(pngs), 4):
            contact_sheet(pngs[i:i + 4], os.path.join(a.outdir, f"hoja-{i // 4 + 1}.png"))
    print(f"{len(pngs)} slides en {a.outdir}")
    if issues:
        print(f"{len(issues)} problema(s):")
        for i in issues:
            print("  ·", i)
        sys.exit(2)
    print("sin desbordes de texto")


if __name__ == "__main__":
    main()
