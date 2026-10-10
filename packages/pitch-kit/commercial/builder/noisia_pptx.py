"""
noisia_pptx.py — the commercial slide library, drawn natively in python-pptx.

Every slide is real PowerPoint: text boxes, rounded rectangles and PNG icons, in Google Sans,
on a 1920x1080 grid. Nothing is a screenshot, so a deck opens in Google Slides and every word,
card and logo can be edited, moved or replaced.

Coordinates are written in canvas pixels (the same 1920x1080 the HTML engine uses) and
converted here: 1px = 6350 EMU, and a CSS px is half a point at this size.

The visual recipe is the one the approved study decks converged on (CANON.md, LAYOUTS.md):
white canvas, one soft cyan/red atmosphere per slide in a different corner, cards in #fafafa
with a 1px #eeeeee border and 16px radius, a teal eyebrow with a real icon, titles at 54px,
one accent word in teal, no shadows, no side-tab borders, no gradients on values.
"""
import math, os, re, uuid
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.normpath(os.path.join(HERE, "..", "assets"))
FONT = os.environ.get("NOISIA_FONT", "Google Sans")

# Tokens (engine/noisia-tokens.css + deck-components.css)
INK = "0a0a0a"; FG = "2b2b2b"; FG2 = "6d6d6d"; EDGE = "8f8f8f"
TEAL = "008a8a"; TEAL_SOFT = "e6f7f7"; CORAL = "e2543c"; CORAL_SOFT = "fdece8"
SURFACE = "fafafa"; BORDER = "eeeeee"; BORDER2 = "e6e6e6"; WHITE = "ffffff"
CYAN = "00eeee"; DARK = "061218"; POS = "008f66"; POS_SOFT = "eaf8f2"

W, H = 1920, 1080
PADX, TOP = 110, 90
CONTENT_W = W - 2 * PADX


def px(v):
    return Emu(int(round(v * 6350)))


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- text fitting
CHAR_EM = 0.52   # average advance of Google Sans, mixed-case Latin text


_MEASURE = None


def _measurer():
    """Real advance widths when Google Sans (or Product Sans) is installed or bundled; the
    average-width estimate otherwise. Measuring is what keeps titles from leaving holes."""
    global _MEASURE
    if _MEASURE is None:
        try:
            from preview import find_fonts, Fonts
            import contextlib, io as _io
            with contextlib.redirect_stdout(_io.StringIO()):
                reg, bold = find_fonts()
            _MEASURE = Fonts(reg, bold) if ("oogle" in reg or "roduct" in reg) else False
        except BaseException:
            _MEASURE = False
    return _MEASURE


def _lines(text, w, size, bold=False):
    m = _measurer()
    n = 0
    for para in str(text).split("\n"):
        words = para.split()
        if m:
            f = m.get(size, bold)
            sp, cur = f.getlength(" "), 0.0
            for wd in words:
                L = f.getlength(wd)
                if cur and cur + sp + L > w:
                    n += 1; cur = L
                else:
                    cur = cur + (sp if cur else 0) + L
            n += 1
            continue
        cpl = max(1, int(w / (size * CHAR_EM)))
        cur = 0
        for wd in words:
            L = len(wd) + (1 if cur else 0)
            if cur + L > cpl:
                n += 1; cur = len(wd)
            else:
                cur += L
        n += 1
    return n


def fit(text, w, h, size, lh=1.3, min_ratio=0.8):
    """Shrink the size until the estimate fits the box. Returns (size, overflow?)."""
    s = size
    while s >= size * min_ratio:
        if _lines(text, w, s) * s * lh <= h:
            return s, False
        s -= 1
    return s, True


class Deck:
    def __init__(self, lang="es", mode="send", header_right="noisia.ai", portrait=False):
        self.lang, self.mode, self.header_right = lang, mode, header_right
        self.prs = Presentation()
        if portrait:   # A4 one-pager
            self.prs.slide_width, self.prs.slide_height = Emu(7560310), Emu(10692130)
        else:
            self.prs.slide_width, self.prs.slide_height = px(W), px(H)
        self._theme_font()
        self.warnings = []
        self.n = 0
        self._bg_cycle = 0

    # ------------------------------------------------------------ plumbing
    def _theme_font(self):
        """New text boxes typed in Google Slides default to the brand font too."""
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        theme = self.prs.slide_master.part.part_related_by(RT.THEME)
        xml = theme.blob.decode("utf-8")
        xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*', r"\g<1>" + FONT, xml)
        theme._blob = xml.encode("utf-8")

    def new(self, bg=None, notes=""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        if bg:
            self.background(s, bg)
        s._noisia_notes = [notes] if notes else []
        return s

    def background(self, s, name):
        path = os.path.join(ASSETS, "bg", name)
        _, rid = s.part.get_or_add_image_part(path)
        bg = etree.fromstring(
            '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/><a:srcRect/>'
            '<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>')
        s.shapes._spTree.getparent().insert(0, bg)

    def content_bg(self):
        self._bg_cycle = self._bg_cycle % 8 + 1
        return f"content-{self._bg_cycle:02d}.png"

    def note(self, s, text):
        if text:
            s._noisia_notes.append(text)

    def finish_notes(self):
        for s in self.prs.slides:
            notes = [n for n in getattr(s, "_noisia_notes", []) if n]
            if notes:
                s.notes_slide.notes_text_frame.text = "\n\n".join(notes)

    def explain(self, s, text):
        """Send version: the sentence that makes the slide readable without a presenter.
        Present version: the same sentence goes to the speaker notes."""
        if not text:
            return None
        if self.mode == "send":
            return text
        self.note(s, text)
        return None

    def detail(self, s, text):
        if not text:
            return None
        if self.mode == "send":
            return text
        self.note(s, "· " + text)
        return None

    # ------------------------------------------------------------ primitives
    def text(self, s, x, y, w, h, content, size=30, color=FG, bold=False, align="l",
             anchor="t", lh=1.3, accent=TEAL, fit_box=True, name=None, italic=False):
        """`**word**` marks the accent run. Returns the shape."""
        content = "" if content is None else str(content)
        if fit_box and content:
            size2, over = fit(re.sub(r"\*\*", "", content), w, h, size, lh)
            if over:
                self.warnings.append(f"slide {self.n}: el texto no cabe ({content[:50]}…)")
            size = size2
        tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        for i, para in enumerate(content.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            p.line_spacing = round(lh / 1.2, 3)   # CSS line-height to PowerPoint 'multiple'
            parts = re.split(r"(\*\*[^*]+\*\*)", para)
            for part in parts:
                if not part:
                    continue
                r = p.add_run()
                is_acc = part.startswith("**")
                r.text = part[2:-2] if is_acc else part
                f = r.font
                f.name = FONT
                f.size = Pt(size * 0.5)
                f.bold = bold
                f.italic = italic
                f.color.rgb = rgb(accent if is_acc else color)
        return tb

    def box(self, s, x, y, w, h, fill=SURFACE, line=BORDER, radius=16, line_w=1, shape="round"):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if shape == "round" else (
            MSO_SHAPE.OVAL if shape == "oval" else MSO_SHAPE.RECTANGLE)
        sh = s.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
        if shape == "round":
            sh.adjustments[0] = min(0.5, radius / max(1, min(w, h)))
        if fill:
            sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgb(line); sh.line.width = Pt(line_w * 0.75)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        sh.text_frame.text = ""
        return sh

    def rule(self, s, x, y, w, color=BORDER, weight=1):
        ln = s.shapes.add_connector(1, px(x), px(y), px(x + w), px(y))
        ln.line.color.rgb = rgb(color); ln.line.width = Pt(weight * 0.75)
        return ln

    def image(self, s, path, x, y, w=None, h=None, name=None):
        pic = s.shapes.add_picture(path, px(x), px(y), px(w) if w else None, px(h) if h else None)
        if name:
            pic.name = name
        return pic

    def icon(self, s, name, x, y, size=28, tone="teal"):
        p = os.path.join(ASSETS, "icons", f"{name}-{tone}.png")
        if not os.path.exists(p):
            self.warnings.append(f"slide {self.n}: ícono '{name}' no existe, usé 'sparks'")
            p = os.path.join(ASSETS, "icons", f"sparks-{tone}.png")
        # the PNG cell carries 8px of padding per 128px, so the glyph lands at `size`
        pad = size * 8 / 112
        return self.image(s, p, x - pad, y - pad, w=size + 2 * pad, h=size + 2 * pad)

    def icobox(self, s, name, x, y, size=56, tone="teal", fill=TEAL_SOFT):
        self.box(s, x, y, size, size, fill=fill, line=None, radius=size * 0.28)
        g = size * 0.5
        self.icon(s, name, x + (size - g) / 2, y + (size - g) / 2, g, tone)

    def brand_chip(self, s, slug, x, y, size=44):
        import json
        colors = json.load(open(os.path.join(ASSETS, "brands", "colors.json")))
        p = os.path.join(ASSETS, "brands", f"{slug}-white.png")
        if slug not in colors or not os.path.exists(p):
            self.warnings.append(f"slide {self.n}: marca '{slug}' no está en el set")
            return
        self.box(s, x, y, size, size, fill=colors[slug].lstrip("#"), line=None, radius=size * 0.26)
        g = size * 0.56
        pad = g * 8 / 112
        self.image(s, p, x + (size - g) / 2 - pad, y + (size - g) / 2 - pad, w=g + 2 * pad, h=g + 2 * pad)

    def pill(self, s, x, y, label, fill=TEAL_SOFT, color=TEAL, size=21, padx=22, h=None, bold=True):
        h = h or size * 2.1
        w = len(label) * size * 0.56 + 2 * padx
        self.box(s, x, y, w, h, fill=fill, line=None, radius=h / 2)
        self.text(s, x, y, w, h, label, size, color, bold=bold, align="c", anchor="m", fit_box=False)
        return w

    def slide_number(self, s, x, y, w, h, dark=False):
        tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        fld = etree.SubElement(p._p, qn("a:fld"))
        fld.set("id", "{" + str(uuid.uuid4()).upper() + "}")
        fld.set("type", "slidenum")
        rpr = etree.SubElement(fld, qn("a:rPr"))
        rpr.set("lang", "es-MX"); rpr.set("sz", "775"); rpr.set("b", "1")
        fill = etree.SubElement(rpr, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", "ffffff" if dark else FG)
        latin = etree.SubElement(rpr, qn("a:latin")); latin.set("typeface", FONT)
        t = etree.SubElement(fld, qn("a:t")); t.text = str(self.n)

    def chrome(self, s, dark=False, header_right=None, footer=True):
        """Header (logo + right label) and footer (tagline + live slide number)."""
        logo = "noisia-white.png" if dark else "noisia.png"
        self.image(s, os.path.join(ASSETS, "logo", logo), PADX, 62, h=30, name="Logo Noisia")
        edge = "c9d3d6" if dark else EDGE
        hr = self.header_right if header_right is None else header_right
        if hr:
            self.text(s, W - PADX - 600, 66, 600, 24, hr, 15.5, edge, align="r", fit_box=False)
        if footer:
            self.text(s, PADX, H - 66, 700, 22, "noisia · social intelligence architects", 15.5, edge,
                      fit_box=False)
            self.slide_number(s, W - PADX - 200, H - 66, 200, 22, dark)

    def eyebrow(self, s, text, icon=None, x=PADX, y=158, color=TEAL, tone="teal"):
        if not text:
            return
        dx = 0
        if icon:
            self.icon(s, icon, x, y + 1, 24, tone)
            dx = 34
        self.text(s, x + dx, y, 1200, 30, text, 21, color, bold=True, fit_box=False)

    def title(self, s, text, y=200, w=1500, size=54, color=INK, h=None):
        h = h or size * 2.3
        return self.text(s, PADX, y, w, h, text, size, color, lh=1.08)

    def head(self, s, d, w=1500, size=54):
        """Eyebrow + title + optional explain line. Returns the y where content can start."""
        self.eyebrow(s, d.get("eyebrow"), d.get("eyebrow_icon"))
        t = d.get("title", "")
        size2, _ = fit(t.replace("**", ""), w, size * 2.2, size, 1.08)
        lines = min(2, _lines(t.replace("**", ""), w, size2))
        self.title(s, t, w=w, size=size2, h=size2 * 1.1 * lines + 6)
        y = 200 + size2 * 1.1 * lines + 18
        ex = self.explain(s, d.get("explain"))
        if ex:
            self.text(s, PADX, y, min(w, 1400), 70, ex, 24, FG2, lh=1.35)
            y += 24 * 1.35 * min(2, _lines(ex, min(w, 1400), 24)) + 8
        return y + 30

    # ================================================================ SLIDES
    def cover(self, d):
        """Kit cover: illustration bleeding right, copy anchored bottom-left. With `client`,
        it becomes the prepared-for cover and carries a client-logo slot."""
        s = self.new(notes=d.get("notes"))
        art = os.path.join(ASSETS, "art", "cover-illustration.png")
        self.image(s, art, 1010, -70, h=1220, name="Ilustración")
        self.chrome(s, header_right=d.get("header_right", self.header_right))
        y = 330 if d.get("client") else 380
        if d.get("eyebrow"):
            self.pill(s, PADX, y, d["eyebrow"])
            y += 74
        size, _ = fit(d["title"].replace("**", ""), 1000, 290, 76, 1.04)
        lines = min(4, _lines(d["title"].replace("**", ""), 1000, size))
        self.text(s, PADX, y, 1000, size * 1.08 * lines + 10, d["title"], size, INK, lh=1.04)
        y += size * 1.08 * lines + 30
        if d.get("subtitle"):
            self.text(s, PADX, y, 900, 80, d["subtitle"], 26, FG2, lh=1.35)
            y += 26 * 1.35 * min(3, _lines(d["subtitle"], 900, 26)) + 22
        if d.get("meta"):
            self.text(s, PADX, y, 1000, 30, d["meta"], 19, FG2, fit_box=False)
            y += 50
        if d.get("client"):
            lab = {"es": "Preparado para", "en": "Prepared for"}[self.lang]
            self.text(s, PADX, y + 10, 400, 26, lab, 17, EDGE, bold=True, fit_box=False)
            logo = d.get("client_logo")
            if logo and os.path.exists(logo):
                self.box(s, PADX, y + 44, 260, 84, fill=WHITE, line=BORDER2, radius=14)
                self.image(s, logo, PADX + 18, y + 54, h=64, name="Logo del cliente")
            else:
                slot = os.path.join(ASSETS, "art", f"client-logo-slot-{self.lang}.png")
                self.image(s, slot, PADX, y + 44, w=260, h=92, name="Logo del cliente (reemplazar)")
                self.text(s, PADX, y + 44, 260, 92, d["client"], 22, FG2, bold=True, align="c",
                          anchor="m", fit_box=False)
                self.note(s, {"es": "Logo del cliente: clic derecho sobre el recuadro > Reemplazar imagen.",
                              "en": "Client logo: right click the box > Replace image."}[self.lang])
        return s

    def statement(self, d):
        """A punch line on the bold brand art. `dark: true` puts it on the dark answer canvas."""
        dark = d.get("dark", False)
        s = self.new(bg=("dark-01.png" if dark else d.get("bg", "bold-right.jpg")), notes=d.get("notes"))
        self.chrome(s, dark=dark)
        col = WHITE if dark else INK
        acc = CYAN if dark else TEAL
        w = d.get("width", 960)
        y = 300
        if d.get("kicker"):
            self.text(s, PADX, y, w, 34, d["kicker"], 24, acc, bold=True, fit_box=False)
            y += 64
        size, _ = fit(d["title"].replace("**", ""), w, 420, d.get("size", 80), 1.04)
        lines = min(5, _lines(d["title"].replace("**", ""), w, size))
        self.text(s, PADX, y, w, size * 1.08 * lines + 10, d["title"], size, col, lh=1.04, accent=acc)
        y += size * 1.08 * lines + 36
        sub = d.get("subtitle")
        if sub:
            self.text(s, PADX, y, min(w, 980), 120, sub, 28, "c9d3d6" if dark else FG2, lh=1.4)
            y += 28 * 1.4 * min(3, _lines(sub, min(w, 980), 28)) + 24
        ex = self.explain(s, d.get("explain"))
        if ex:
            self.text(s, PADX, y, min(w, 980), 110, ex, 22, "c9d3d6" if dark else FG2, lh=1.4)
        return s

    def questions(self, d):
        """Questions this brand could answer with us. Researched from public facts about the
        brand and its industry; never social-data insights (COMMERCIAL.md, rule 1)."""
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        if d.get("context"):
            self.text(s, PADX, y - 10, 1500, 60, d["context"], 24, FG, lh=1.35)
            y += 24 * 1.35 * min(2, _lines(d["context"], 1500, 24)) + 16
        items = d["items"][:5]
        bottom = 876 if d.get("answers") else 960
        rowh = (bottom - y) / len(items)
        for i, it in enumerate(items):
            ry = y + i * rowh
            self.rule(s, PADX, ry, CONTENT_W)
            self.icobox(s, it.get("icon", "help-circle"), PADX, ry + 18, 48)
            self.text(s, PADX + 72, ry + 20, 300, 30, it.get("theme", ""), 18, TEAL, bold=True, fit_box=False)
            det = self.detail(s, it.get("detail"))
            qx, qw = PADX + 400, CONTENT_W - 400
            if det:
                self.text(s, qx, ry + 16, qw, rowh * 0.46, it["question"], 25, INK, lh=1.2)
                self.text(s, qx, ry + 16 + rowh * 0.46, qw, rowh * 0.44, det, 19, FG2, lh=1.3)
            else:
                self.text(s, qx, ry + 16, qw, rowh - 28, it["question"], 27, INK, lh=1.2)
        if d.get("answers"):
            ax = PADX
            aw = (CONTENT_W - 2 * 20) / 3
            for i, a in enumerate(d["answers"][:3]):
                bx = ax + i * (aw + 20)
                self.box(s, bx, 896, aw, 70, fill=TEAL_SOFT if i == 0 else SURFACE,
                         line=None if i == 0 else BORDER, radius=14)
                self.text(s, bx + 22, 896, 130, 70, a["name"], 20, TEAL, bold=True, anchor="m", fit_box=False)
                self.text(s, bx + 150, 900, aw - 170, 62, a["text"], 18, FG, anchor="m", lh=1.25)
        return s

    def what_we_do(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        self.eyebrow(s, d.get("eyebrow"), d.get("eyebrow_icon"))
        tsize, _ = fit(d["title"].replace("**", ""), 760, 330, 58, 1.08)
        tl = min(5, _lines(d["title"].replace("**", ""), 760, tsize))
        self.text(s, PADX, 210, 760, tsize * 1.1 * tl + 8, d["title"], tsize, INK, lh=1.08)
        if d.get("lead"):
            self.text(s, PADX, 210 + tsize * 1.1 * tl + 40, 720, 300, d["lead"], 27, FG2, lh=1.42)
        x0, wcol = 1000, 810
        items = d["pillars"][:4]
        y0, gap = 230, 24
        hh = (930 - y0 - gap * (len(items) - 1)) / len(items)
        for i, it in enumerate(items):
            by = y0 + i * (hh + gap)
            first = i == d.get("highlight", -1)
            self.box(s, x0, by, wcol, hh, fill=TEAL_SOFT if first else SURFACE,
                     line=None if first else BORDER, radius=18)
            self.icobox(s, it.get("icon", "sparks"), x0 + 30, by + 30, 56, fill=WHITE if first else TEAL_SOFT)
            self.text(s, x0 + 112, by + 28, wcol - 140, 40, it["name"], 27, INK, bold=True)
            det = self.detail(s, it.get("text")) if d.get("text_is_detail") else it.get("text")
            if det:
                self.text(s, x0 + 112, by + 72, wcol - 140, hh - 90, det, 21, FG2, lh=1.35)
        return s

    def where_we_stand(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        cols = d["columns"]
        gap = 24
        weights = [1.25 if c.get("noisia") else 1 for c in cols]
        unit = (CONTENT_W - gap * (len(cols) - 1)) / sum(weights)
        x = PADX
        top, hh = y, 960 - y
        for c, wt in zip(cols, weights):
            cw = unit * wt
            hi = c.get("noisia")
            self.box(s, x, top, cw, hh, fill=TEAL_SOFT if hi else SURFACE, line=None if hi else BORDER, radius=20)
            if hi:
                self.image(s, os.path.join(ASSETS, "logo", "noisia.png"), x + 36, top + 40, h=36)
            else:
                self.text(s, x + 36, top + 38, cw - 72, 44, c["name"], 28, INK, bold=True)
            if c.get("kicker"):
                self.text(s, x + 36, top + 96, cw - 72, 60, c["kicker"], 21, TEAL if hi else FG2, bold=True, lh=1.25)
            py = top + 190
            pts = c.get("points", [])
            for ptxt in pts:
                ic = "check-circle" if hi else "xmark-circle"
                self.icon(s, ic, x + 36, py + 4, 28, "teal" if hi else "grey")
                self.text(s, x + 80, py, cw - 116, 80, ptxt, 25, FG if hi else FG2, lh=1.3)
                py += 25 * 1.3 * min(3, _lines(ptxt, cw - 116, 25)) + 30
            if c.get("verdict"):
                self.rule(s, x + 36, top + hh - 140, cw - 72, BORDER2 if not hi else "bfe6e6")
                self.text(s, x + 36, top + hh - 122, cw - 72, 100, c["verdict"], 23,
                          TEAL if hi else FG, bold=bool(hi), lh=1.35)
            x += cw + gap
        return s

    def coverage(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        groups = d["groups"][:3]
        gap = 24
        cw = (CONTENT_W - gap * (len(groups) - 1)) / len(groups)
        hh = min(420, 760 - y)
        for i, g in enumerate(groups):
            x = PADX + i * (cw + gap)
            self.box(s, x, y, cw, hh, radius=20)
            bx = x + 34
            for slug in g.get("brands", [])[:7]:
                self.brand_chip(s, slug, bx, y + 34, 46)
                bx += 56
            if g.get("more"):
                self.pill(s, bx + 4, y + 40, g["more"], fill=WHITE, color=FG2, size=17, padx=14, h=34)
            if g.get("icon") and not g.get("brands"):
                self.icobox(s, g["icon"], x + 34, y + 34, 50)
            self.text(s, x + 34, y + 116, cw - 68, 44, g["name"], 30, INK, bold=True)
            self.text(s, x + 34, y + 170, cw - 68, hh - 190, g["text"], 23, FG2, lh=1.4)
        stats = d.get("stats", [])
        if stats:
            sy = y + hh + 50
            self.rule(s, PADX, sy, CONTENT_W)
            sw = CONTENT_W / len(stats)
            for i, st in enumerate(stats):
                x = PADX + i * sw
                self.text(s, x, sy + 34, sw, 96, st["value"], 72, TEAL, bold=True, fit_box=False)
                self.text(s, x, sy + 130, sw - 30, 40, st["label"], 23, FG2)
        if d.get("foot"):
            self.text(s, PADX, 960, CONTENT_W, 30, d["foot"], 18, EDGE, italic=True, fit_box=False)
        return s

    def portfolio(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        gap = 28
        cw = (CONTENT_W - gap) / 2
        bottom = 880 if d.get("punchline") else 960
        hh = bottom - y
        for i, side in enumerate([d["left"], d["right"]]):
            x = PADX + i * (cw + gap)
            acc = TEAL if i == 0 else CORAL
            soft = TEAL_SOFT if i == 0 else CORAL_SOFT
            self.box(s, x, y, cw, hh, radius=22)
            self.pill(s, x + 36, y + 32, side["kicker"], fill=soft, color=acc, size=17, padx=16, h=34)
            self.text(s, x + 36, y + 84, cw - 72, 60, side["name"], 44, INK, bold=False)
            self.text(s, x + 36, y + 146, cw - 72, 64, side["lead"], 21, FG2, lh=1.35)
            qy = y + 222
            qs = side.get("questions", [])[:5]
            qh = (hh - 222 - 86) / max(1, len(qs))
            for q in qs:
                self.icon(s, "help-circle", x + 36, qy + 4, 24, "teal" if i == 0 else "coral")
                self.text(s, x + 74, qy, cw - 110, qh, q, 22, FG, lh=1.28)
                qy += qh
            tx = x + 36
            for t in side.get("tags", [])[:5]:
                tw = self.pill(s, tx, y + hh - 62, t, fill=WHITE, color=FG2, size=16, padx=14, h=34, bold=False)
                tx += tw + 10
        if d.get("punchline"):
            self.text(s, PADX, 906, CONTENT_W, 50, d["punchline"], 30, INK, align="c", accent=TEAL)
        return s

    def outputs_impact(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        lw = 980
        self.text(s, PADX, y, lw, 30, d.get("outputs_label", ""), 20, TEAL, bold=True, fit_box=False)
        outs = d["outputs"][:4]
        gap = 20
        cw = (lw - gap) / 2
        ch = (960 - y - 50 - gap) / 2
        for i, o in enumerate(outs):
            r, c = divmod(i, 2)
            x, yy = PADX + c * (cw + gap), y + 46 + r * (ch + gap)
            self.box(s, x, yy, cw, ch, radius=18)
            self.icobox(s, o.get("icon", "reports"), x + 28, yy + 28, 52)
            self.text(s, x + 28, yy + 100, cw - 56, 36, o["name"], 25, INK, bold=True)
            self.text(s, x + 28, yy + 142, cw - 56, ch - 160, o["text"], 20, FG2, lh=1.35)
        rx = PADX + lw + 50
        rw = CONTENT_W - lw - 50
        self.box(s, rx, y, rw, 960 - y, fill=TEAL_SOFT, line=None, radius=22)
        self.text(s, rx + 40, y + 36, rw - 80, 30, d.get("impact_label", ""), 20, TEAL, bold=True, fit_box=False)
        iy = y + 90
        imps = d["impacts"][:4]
        ih = (960 - y - 120) / max(1, len(imps))
        for im in imps:
            self.icon(s, im.get("icon", "check-circle"), rx + 40, iy + 4, 28)
            self.text(s, rx + 84, iy, rw - 124, ih - 10, im["text"], 23, INK, lh=1.3)
            iy += ih
        return s

    def cooperations(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        items = d["items"][:3]
        gap = 24
        cw = (CONTENT_W - gap * (len(items) - 1)) / len(items)
        bottom = 880 if d.get("foot") else 960
        hh = bottom - y
        for i, it in enumerate(items):
            x = PADX + i * (cw + gap)
            hi = i == d.get("highlight", -1)
            self.box(s, x, y, cw, hh, fill=TEAL_SOFT if hi else SURFACE, line=None if hi else BORDER, radius=20)
            self.icobox(s, it.get("icon", "sparks"), x + 34, y + 34, 56, fill=WHITE if hi else TEAL_SOFT)
            if it.get("kicker"):
                self.text(s, x + 34, y + 112, cw - 68, 28, it["kicker"], 17, TEAL, bold=True, fit_box=False)
            self.text(s, x + 34, y + 146, cw - 68, 90, it["name"], 30, INK, bold=True, lh=1.15)
            self.text(s, x + 34, y + 250, cw - 68, hh - 380, it["text"], 21, FG2, lh=1.38)
            if it.get("price"):
                self.rule(s, x + 34, y + hh - 110, cw - 68, BORDER2)
                self.text(s, x + 34, y + hh - 92, cw - 68, 26, it.get("price_label", ""), 17, FG2, fit_box=False)
                self.text(s, x + 34, y + hh - 62, cw - 68, 46, it["price"], 34, TEAL, bold=True, fit_box=False)
        if d.get("foot"):
            self.box(s, PADX, 900, CONTENT_W, 60, fill=TEAL_SOFT, line=None, radius=14)
            self.text(s, PADX + 30, 900, CONTENT_W - 60, 60, d["foot"], 22, INK, anchor="m", accent=TEAL)
        return s

    def engine(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        hh = 960 - y
        lw, mw, gap = 470, 640, 60
        rw = CONTENT_W - lw - mw - 2 * gap
        # sources
        self.box(s, PADX, y, lw, hh, radius=20)
        self.text(s, PADX + 34, y + 30, lw - 68, 28, d["sources_label"], 20, TEAL, bold=True, fit_box=False)
        sy = y + 80
        for src in d["sources"][:5]:
            bx = PADX + 34
            for slug in src.get("brands", [])[:5]:
                self.brand_chip(s, slug, bx, sy, 34)
                bx += 42
            if src.get("icon") and not src.get("brands"):
                self.icobox(s, src["icon"], bx, sy, 34)
                bx += 42
            self.text(s, PADX + 34, sy + 42, lw - 68, 50, src["name"], 20, FG, bold=True, lh=1.25)
            sy += (hh - 110) / max(1, len(d["sources"][:5]))
        # arrows
        ay = y + hh / 2 - 18
        self.icon(s, "arrow-right", PADX + lw + gap / 2 - 18, ay, 36, "grey")
        mx = PADX + lw + gap
        self.icon(s, "arrow-right", mx + mw + gap / 2 - 18, ay, 36, "grey")
        # brain
        self.box(s, mx, y, mw, hh, fill=TEAL_SOFT, line=None, radius=22)
        self.image(s, os.path.join(ASSETS, "logo", "noisia.png"), mx + 40, y + 36, h=34)
        self.text(s, mx + 40, y + 92, mw - 80, 70, d["brain_title"], 30, INK, bold=True, lh=1.15)
        self.text(s, mx + 40, y + 170, mw - 80, 80, d.get("brain_text", ""), 20, FG2, lh=1.35)
        my = y + 270
        chips = d.get("methods", [])
        cx = mx + 40
        for m in chips:
            cwid = len(m) * 18 * 0.56 + 32
            if cx + cwid > mx + mw - 40:
                cx = mx + 40; my += 52
            self.pill(s, cx, my, m, fill=WHITE, color=TEAL, size=18, padx=16, h=40)
            cx += cwid + 10
        # impact
        rx = mx + mw + gap
        self.box(s, rx, y, rw, hh, radius=20)
        self.text(s, rx + 34, y + 30, rw - 68, 28, d["impact_label"], 20, TEAL, bold=True, fit_box=False)
        iy = y + 84
        imps = d["impacts"][:4]
        ih = (hh - 110) / max(1, len(imps))
        for im in imps:
            self.icon(s, im.get("icon", "check-circle"), rx + 34, iy + 2, 26)
            self.text(s, rx + 72, iy, rw - 106, ih - 8, im["text"], 21, INK, lh=1.3)
            iy += ih
        return s

    def steps(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        st = d["steps"][:5]
        arrow = 46
        cw = (CONTENT_W - arrow * (len(st) - 1)) / len(st)
        bottom = 820 if d.get("note") else 960
        hh = min(bottom - y, 400)
        for i, it in enumerate(st):
            x = PADX + i * (cw + arrow)
            hi = i == 0
            self.box(s, x, y, cw, hh, fill=TEAL_SOFT if hi else SURFACE, line=None if hi else BORDER, radius=20)
            self.text(s, x + 32, y + 30, 120, 50, f"{i+1:02d}", 40, TEAL, bold=True, fit_box=False)
            if it.get("tag"):
                tw = len(it["tag"]) * 16 * 0.6 + 28
                self.pill(s, x + cw - 32 - tw, y + 36, it["tag"], fill=WHITE if hi else TEAL_SOFT,
                          color=TEAL, size=16, padx=14, h=34)
            self.text(s, x + 32, y + 104, cw - 64, 84, it["name"], 32, INK, bold=True, lh=1.15)
            self.text(s, x + 32, y + 196, cw - 64, hh - 220, it["text"], 23, FG2, lh=1.4)
            if i < len(st) - 1:
                self.icon(s, "arrow-right", x + cw + (arrow - 30) / 2, y + hh / 2 - 15, 30, "grey")
        if d.get("note"):
            ny = y + hh + 30
            self.box(s, PADX, ny, CONTENT_W, 100, fill=TEAL_SOFT, line=None, radius=16)
            self.icon(s, "light-bulb", PADX + 32, ny + 35, 30)
            self.text(s, PADX + 84, ny, CONTENT_W - 120, 100, d["note"], 24, INK, anchor="m", lh=1.35)
        return s

    def catalog(self, d):
        """Study catalogue. With prices: a table (name | what you get | price). Without: a grid."""
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        items = d["items"]
        priced = any(it.get("price") for it in items)
        bottom = 920 if d.get("foot") else 960
        if priced:
            hdr = d.get("headers", ["", "", ""])
            self.text(s, PADX, y, 500, 26, hdr[0], 16, EDGE, bold=True, fit_box=False)
            self.text(s, PADX + 520, y, 900, 26, hdr[1], 16, EDGE, bold=True, fit_box=False)
            self.text(s, W - PADX - 260, y, 260, 26, hdr[2], 16, EDGE, bold=True, align="r", fit_box=False)
            y += 36
            rh = (bottom - y) / len(items)
            for i, it in enumerate(items):
                ry = y + i * rh
                self.rule(s, PADX, ry, CONTENT_W)
                self.text(s, PADX, ry, 500, rh, it["name"], 23, INK, bold=True, anchor="m")
                self.text(s, PADX + 520, ry + 4, 940, rh - 8, "\n".join(it.get("bullets", [])[:2]), 18, FG2,
                          anchor="m", lh=1.25)
                self.text(s, W - PADX - 260, ry, 260, rh, it.get("price", ""), 26, TEAL, bold=True,
                          align="r", anchor="m", fit_box=False)
        else:
            cols = 3
            rows = math.ceil(len(items) / cols)
            gap = 18
            cw = (CONTENT_W - gap * (cols - 1)) / cols
            ch = (bottom - y - gap * (rows - 1)) / rows
            for i, it in enumerate(items):
                r, c = divmod(i, cols)
                x, yy = PADX + c * (cw + gap), y + r * (ch + gap)
                self.box(s, x, yy, cw, ch, radius=16)
                self.icon(s, it.get("icon", "sparks"), x + 28, yy + 28, 30)
                self.text(s, x + 72, yy + 24, cw - 98, 40, it["name"], 25, INK, bold=True)
                by = yy + 82
                for b in it.get("bullets", [])[:2]:
                    n = min(2, _lines(b, cw - 56, 20))
                    self.text(s, x + 28, by, cw - 56, 20 * 1.3 * n + 4, b, 20, FG2, lh=1.3)
                    by += 20 * 1.3 * n + 10
        if d.get("foot"):
            self.text(s, PADX, 935, CONTENT_W, 30, d["foot"], 18, EDGE, italic=True)
        return s

    def detail_slide(self, d):
        """Report or study detail: what it answers on the left, how it is delivered on the right."""
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d, w=1000)
        lw = 1000
        if d.get("lead"):
            self.text(s, PADX, y - 6, lw - 60, 110, d["lead"], 24, FG, lh=1.4)
            y += 24 * 1.4 * min(3, _lines(d["lead"], lw - 60, 24)) + 24
        self.text(s, PADX, y, lw, 28, d.get("list_label", ""), 20, TEAL, bold=True, fit_box=False)
        ly = y + 44
        lst = d.get("list", [])[:6]
        for it in lst:
            self.icon(s, "check-circle", PADX, ly + 3, 26)
            n = min(3, _lines(it, lw - 100, 23))
            self.text(s, PADX + 42, ly, lw - 100, 23 * 1.32 * n + 6, it, 23, FG, lh=1.32)
            ly += 23 * 1.32 * n + 30
        rx = PADX + lw + 40
        rw = CONTENT_W - lw - 40
        self.box(s, rx, 200, rw, 760, radius=22)
        sy = 236
        for blk in d.get("side", [])[:3]:
            self.text(s, rx + 36, sy, rw - 72, 28, blk["label"], 18, TEAL, bold=True, fit_box=False)
            sy += 42
            if blk.get("chips"):
                cx = rx + 36
                for c in blk["chips"]:
                    cwid = len(c) * 17 * 0.6 + 30
                    if cx + cwid > rx + rw - 36:
                        cx = rx + 36; sy += 48
                    self.pill(s, cx, sy, c, fill=WHITE, color=FG, size=17, padx=15, h=38, bold=False)
                    cx += cwid + 10
                sy += 66
            for item in blk.get("items", []):
                self.icon(s, blk.get("icon", "nav-arrow-right"), rx + 36, sy + 2, 22, "teal")
                self.text(s, rx + 68, sy, rw - 104, 60, item, 20, FG, lh=1.3)
                sy += 20 * 1.3 * min(2, _lines(item, rw - 104, 20)) + 14
            sy += 18
        if d.get("sample"):
            self.box(s, rx + 36, 960 - 36 - 80, rw - 72, 80, fill=WHITE, line=BORDER2, radius=14)
            self.icon(s, "arrow-up-right", rx + 60, 960 - 36 - 80 + 26, 28)
            self.text(s, rx + 104, 960 - 36 - 80, rw - 160, 80, d["sample"], 19, INK, bold=True, anchor="m")
        return s

    def factor_pricing(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        lw = 820
        self.box(s, PADX, y, lw, 960 - y, radius=22)
        hdr = d["headers"]
        cx = [PADX + 40, PADX + 260, PADX + 560]
        for j, h in enumerate(hdr):
            self.text(s, cx[j], y + 34, 280, 26, h, 17, EDGE, bold=True, fit_box=False)
        rows = d["rows"]
        ry = y + 78
        rh = (960 - y - 78 - 110) / len(rows)
        for r in rows:
            self.rule(s, PADX + 40, ry, lw - 80)
            self.text(s, cx[0], ry, 200, rh, r[0], 30, INK, bold=True, anchor="m", fit_box=False)
            self.text(s, cx[1], ry, 280, rh, r[1], 28, TEAL, bold=True, anchor="m", fit_box=False)
            self.text(s, cx[2], ry, 240, rh, r[2], 24, FG2, anchor="m", fit_box=False)
            ry += rh
        if d.get("table_foot"):
            self.text(s, PADX + 40, 960 - 90, lw - 80, 60, d["table_foot"], 18, FG2, italic=True, lh=1.3)
        rx = PADX + lw + 50
        rw = CONTENT_W - lw - 50
        feats = d["features"][:5]
        fh = (960 - y) / len(feats)
        for i, f in enumerate(feats):
            fy = y + i * fh
            if i:
                self.rule(s, rx, fy, rw)
            self.icon(s, f.get("icon", "check-circle"), rx, fy + 22, 28)
            self.text(s, rx + 46, fy + 18, rw - 46, 32, f["name"], 22, INK, bold=True)
            self.text(s, rx + 46, fy + 54, rw - 46, fh - 64, f["text"], 18, FG2, lh=1.3)
        return s

    def build_vs_buy(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        gap = 70
        cw = (CONTENT_W - gap) / 2
        bottom = 870 if d.get("foot") else 960
        self.text(s, PADX, y, cw, 32, d["left_label"], 22, FG2, bold=True, fit_box=False)
        self.box(s, PADX + cw + gap, y - 14, cw, bottom - y + 14, fill=TEAL_SOFT, line=None, radius=20)
        self.image(s, os.path.join(ASSETS, "logo", "noisia.png"), PADX + cw + gap + 36, y, h=30)
        rows = d["rows"]
        ry = y + 56
        rh = (bottom - ry) / len(rows)
        for a, b in rows:
            self.rule(s, PADX, ry, cw, BORDER2)
            self.icon(s, "xmark-circle", PADX, ry + rh / 2 - 12, 24, "grey")
            self.text(s, PADX + 40, ry, cw - 40, rh, a, 21, FG2, anchor="m", lh=1.3)
            self.icon(s, "arrow-right", PADX + cw + gap / 2 - 14, ry + rh / 2 - 14, 28, "grey")
            self.icon(s, "check-circle", PADX + cw + gap + 36, ry + rh / 2 - 12, 24)
            self.text(s, PADX + cw + gap + 76, ry, cw - 112, rh, b, 21, INK, anchor="m", lh=1.3)
            ry += rh
        if d.get("foot"):
            self.text(s, PADX, 895, CONTENT_W, 60, d["foot"], 24, INK, lh=1.35, accent=TEAL)
        return s

    def year_plan(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        lw = 470
        tx = PADX + lw + 30
        tw = CONTENT_W - lw - 30
        mw = tw / 12
        plans = d["plans"][:3]
        ph = (960 - y) / len(plans)
        for i, p in enumerate(plans):
            py = y + i * ph
            if i:
                self.rule(s, PADX, py, CONTENT_W)
            self.text(s, PADX, py + 20, 50, 48, f"{i+1}", 34, TEAL, bold=True, fit_box=False)
            self.text(s, PADX + 50, py + 24, lw - 60, 34, p["name"], 23, INK, bold=True)
            self.text(s, PADX + 50, py + 64, lw - 60, ph - 80, p["text"], 18, FG2, lh=1.3)
            ty = py + 28
            for m in range(12):
                self.text(s, tx + m * mw, ty, mw, 22, str(m + 1), 14, EDGE, align="c", fit_box=False)
            self.rule(s, tx, ty + 30, tw, BORDER2)
            for blk in p["blocks"]:
                bx = tx + (blk["start"] - 1) * mw + 3
                bw = blk["span"] * mw - 6
                first = blk.get("first", False)
                self.box(s, bx, ty + 44, bw, ph - 96, fill=TEAL_SOFT if first else SURFACE,
                         line=None if first else BORDER, radius=12)
                self.text(s, bx + 16, ty + 54, bw - 32, 30, blk["name"], 19, INK, bold=True)
                if blk.get("price"):
                    self.text(s, bx + 16, ty + 86, bw - 32, 30, blk["price"], 19, TEAL, bold=True)
        return s

    def understood(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        lw = 820
        self.box(s, PADX, y, lw, 960 - y, fill=TEAL_SOFT, line=None, radius=22)
        self.text(s, PADX + 44, y + 40, lw - 88, 28, d.get("question_label", ""), 19, TEAL, bold=True, fit_box=False)
        self.text(s, PADX + 44, y + 96, lw - 88, 960 - y - 140, d["question"], 46, INK, lh=1.16)
        rx = PADX + lw + 50
        rw = CONTENT_W - lw - 50
        self.text(s, rx, y, rw, 28, d.get("decisions_label", ""), 19, TEAL, bold=True, fit_box=False)
        decs = d["decisions"][:4]
        dh = (960 - y - 50) / len(decs)
        for i, dec in enumerate(decs):
            dy = y + 50 + i * dh
            self.rule(s, rx, dy, rw)
            self.text(s, rx, dy + 18, 70, 50, f"{i+1:02d}", 34, TEAL, bold=True, fit_box=False)
            if isinstance(dec, dict):
                self.text(s, rx + 80, dy + 20, rw - 80, 36, dec["name"], 24, INK, bold=True)
                det = self.detail(s, dec.get("text"))
                if det:
                    self.text(s, rx + 80, dy + 60, rw - 80, dh - 70, det, 19, FG2, lh=1.3)
            else:
                self.text(s, rx + 80, dy + 16, rw - 80, dh - 30, dec, 24, INK, lh=1.3)
        return s

    def formats(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        items = d["items"][:3]
        gap = 24
        cw = (CONTENT_W - gap * (len(items) - 1)) / len(items)
        bottom = 900 if d.get("foot") else 960
        need = max(150 + 20 * 1.36 * _lines(it.get("text", ""), cw - 68, 20) + 40 + 160 for it in items)
        hh = min(bottom - y, max(440, need))
        for i, it in enumerate(items):
            x = PADX + i * (cw + gap)
            hi = i == d.get("highlight", 0)
            self.box(s, x, y, cw, hh, fill=TEAL_SOFT if hi else SURFACE, line=None if hi else BORDER, radius=20)
            self.text(s, x + 34, y + 30, 80, 44, f"{i+1:02d}", 34, TEAL, bold=True, fit_box=False)
            if it.get("kicker"):
                self.text(s, x + 120, y + 40, cw - 154, 26, it["kicker"], 16, TEAL, bold=True, fit_box=False)
            self.text(s, x + 34, y + 96, cw - 68, 50, it["name"], 32, INK, bold=True)
            self.text(s, x + 34, y + 150, cw - 68, hh - 330, it["text"], 20, FG2, lh=1.36)
            my = y + hh - 160
            for m in it.get("meta", [])[:2]:
                self.rule(s, x + 34, my, cw - 68, "bfe6e6" if hi else BORDER2)
                self.text(s, x + 34, my + 12, cw - 68, 22, m["label"], 15, EDGE, bold=True, fit_box=False)
                self.text(s, x + 34, my + 38, cw - 68, 30, m["value"], 19, INK, bold=True)
                my += 78
        if d.get("foot"):
            self.text(s, PADX, y + hh + 22, CONTENT_W, 40, d["foot"], 20, FG2, italic=True)
        return s

    def blocks(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        items = d["items"][:8]
        cols = 4 if len(items) > 4 else len(items)
        rows = math.ceil(len(items) / cols)
        gap = 18
        cw = (CONTENT_W - gap * (cols - 1)) / cols
        ch = min((960 - y - gap * (rows - 1)) / rows, 330)
        for i, it in enumerate(items):
            r, c = divmod(i, cols)
            x, yy = PADX + c * (cw + gap), y + r * (ch + gap)
            self.box(s, x, yy, cw, ch, radius=16)
            self.text(s, x + 26, yy + 22, 60, 48, it.get("letter", chr(65 + i)), 36, TEAL, bold=True, fit_box=False)
            self.text(s, x + 82, yy + 28, cw - 108, 64, it["name"], 25, INK, bold=True, lh=1.2)
            qy = yy + 100
            for q in it.get("questions", [])[:3]:
                n = min(3, _lines(q, cw - 52, 20))
                self.text(s, x + 26, qy, cw - 52, 20 * 1.32 * n + 6, q, 20, FG2, lh=1.32)
                qy += 20 * 1.32 * n + 12
        return s

    def investment(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d)
        lw = 1060
        rows = d["rows"][:4]
        bottom = 900 if d.get("terms") else 960
        rh = min((bottom - y) / len(rows), 230)
        bottom = y + rh * len(rows)
        for i, r in enumerate(rows):
            ry = y + i * rh
            self.box(s, PADX, ry + 6, lw, rh - 12, fill=TEAL_SOFT if i == 0 else SURFACE,
                     line=None if i == 0 else BORDER, radius=18)
            self.text(s, PADX + 34, ry + 30, lw - 420, 40, r["name"], 30, INK, bold=True)
            self.text(s, PADX + 34, ry + 80, lw - 420, rh - 100, r.get("text", ""), 22, FG2, lh=1.35)
            self.text(s, PADX + lw - 360, ry + 22, 326, 50, r["price"], 38, TEAL, bold=True, align="r", fit_box=False)
            if r.get("unit"):
                self.text(s, PADX + lw - 360, ry + 76, 326, 26, r["unit"], 17, FG2, align="r", fit_box=False)
        rx = PADX + lw + 40
        rw = CONTENT_W - lw - 40
        self.box(s, rx, y + 6, rw, bottom - y - 12, radius=18)
        self.text(s, rx + 34, y + 34, rw - 68, 28, d.get("excluded_label", ""), 19, CORAL, bold=True, fit_box=False)
        ey = y + 84
        for e in d.get("excluded", [])[:7]:
            self.icon(s, "minus", rx + 34, ey + 3, 22, "coral")
            self.text(s, rx + 68, ey, rw - 100, 90, e, 21, FG, lh=1.3)
            ey += 21 * 1.3 * min(3, _lines(e, rw - 100, 21)) + 18
        if d.get("terms"):
            self.text(s, PADX, bottom + 22, CONTENT_W, 40, d["terms"], 18, FG2, italic=True)
        return s

    def sample(self, d):
        s = self.new(bg=self.content_bg(), notes=d.get("notes"))
        self.chrome(s)
        y = self.head(s, d, w=900)
        lw = 860
        if d.get("kicker"):
            self.pill(s, PADX, y, d["kicker"], size=18, padx=16, h=38)
            y += 60
        self.text(s, PADX, y, lw, 110, d["study_title"], 40, INK, lh=1.12)
        y += 40 * 1.12 * min(2, _lines(d["study_title"], lw, 40)) + 24
        self.text(s, PADX, y, lw - 40, 240, d["text"], 25, FG2, lh=1.42)
        y += 25 * 1.42 * min(6, _lines(d["text"], lw - 40, 25)) + 40
        if d.get("link_label"):
            self.box(s, PADX, y, 520, 70, fill=WHITE, line=BORDER2, radius=14)
            self.icon(s, "arrow-up-right", PADX + 24, y + 21, 28)
            self.text(s, PADX + 66, y, 440, 70, d["link_label"], 20, INK, bold=True, anchor="m")
            if d.get("link_url"):
                self.note(s, d["link_url"])
        rx = PADX + lw + 60
        rw = CONTENT_W - lw - 60
        thumb = d.get("thumbnail")
        if thumb and os.path.exists(thumb):
            self.image(s, thumb, rx, 200, w=rw, name="Portada del estudio")
        else:
            self.box(s, rx, 200, rw, rw * 9 / 16, fill=SURFACE, line=BORDER2, radius=18)
            self.text(s, rx, 200, rw, rw * 9 / 16, {"es": "Portada del estudio de muestra",
                                                    "en": "Sample study cover"}[self.lang],
                      20, EDGE, align="c", anchor="m", fit_box=False)
        return s

    def closing(self, d):
        s = self.new(bg=d.get("bg", "bold-corner.jpg"), notes=d.get("notes"))
        self.chrome(s)
        size, _ = fit(d["title"].replace("**", ""), 1000, 360, 72, 1.05)
        lines = min(4, _lines(d["title"].replace("**", ""), 1000, size))
        y = 330
        self.text(s, PADX, y, 1000, size * 1.1 * lines + 10, d["title"], size, INK, lh=1.05)
        y += size * 1.1 * lines + 30
        if d.get("subtitle"):
            self.text(s, PADX, y, 900, 110, d["subtitle"], 27, FG2, lh=1.4)
            y += 27 * 1.4 * min(3, _lines(d["subtitle"], 900, 27)) + 36
        if d.get("cta"):
            w = len(d["cta"]) * 24 * 0.55 + 70
            self.box(s, PADX, y, w, 70, fill=INK, line=None, radius=35)
            self.text(s, PADX, y, w, 70, d["cta"], 24, WHITE, bold=True, align="c", anchor="m", fit_box=False)
            y += 110
        if d.get("contact"):
            self.text(s, PADX, y, 1100, 40, d["contact"], 21, FG2, fit_box=False)
        return s

    def onepager(self, d):
        """A4 portrait: the offer or pilot on one page. Canvas 1191 x 1684 px."""
        PW, PH, M = 1191, 1684, 84
        cw = PW - 2 * M
        s = self.new(notes=d.get("notes"))
        bg = os.path.join(ASSETS, "bg", "content-01.png")
        self.image(s, bg, 0, 0, w=PW, h=PH, name="Fondo")
        self.image(s, os.path.join(ASSETS, "logo", "noisia.png"), M, 64, h=28, name="Logo Noisia")
        logo = d.get("client_logo")
        if logo and os.path.exists(logo):
            self.image(s, logo, PW - M - 200, 54, h=48, name="Logo del cliente")
        elif d.get("client_logo_slot", True):
            slot = os.path.join(ASSETS, "art", f"client-logo-slot-{self.lang}.png")
            self.image(s, slot, PW - M - 180, 52, w=180, h=56, name="Logo del cliente (reemplazar)")
            self.text(s, PW - M - 180, 52, 180, 56, {"es": "Logo del cliente", "en": "Client logo"}[self.lang],
                      14, EDGE, align="c", anchor="m", fit_box=False)
        y = 150
        self.pill(s, M, y, d["label"], size=17, padx=16, h=36)
        y += 62
        tsize, _ = fit(d["title"].replace("**", ""), cw, 160, 46, 1.1)
        tl = min(3, _lines(d["title"].replace("**", ""), cw, tsize))
        self.text(s, M, y, cw, tsize * 1.12 * tl + 6, d["title"], tsize, INK, lh=1.1)
        y += tsize * 1.12 * tl + 22
        self.text(s, M, y, cw, 130, d["intro"], 21, FG2, lh=1.42)
        y += 21 * 1.42 * min(4, _lines(d["intro"], cw, 21)) + 34
        # challenge | decisions
        gap = 26
        colw = (cw - gap) / 2
        boxh = 380
        for i, side in enumerate([d["left"], d["right"]]):
            x = M + i * (colw + gap)
            hi = i == 1
            self.box(s, x, y, colw, boxh, fill=TEAL_SOFT if hi else SURFACE, line=None if hi else BORDER, radius=16)
            self.text(s, x + 26, y + 22, colw - 52, 28, side["label"], 18, TEAL, bold=True, fit_box=False)
            iy = y + 64
            for it in side["items"][:6]:
                self.icon(s, "check-circle" if hi else "nav-arrow-right", x + 26, iy + 2, 20, "teal")
                self.text(s, x + 56, iy, colw - 82, 80, it, 19, FG, lh=1.32)
                iy += 19 * 1.32 * min(3, _lines(it, colw - 82, 19)) + 14
        y += boxh + 30
        # proposal
        self.text(s, M, y, cw, 26, d["proposal_label"], 18, TEAL, bold=True, fit_box=False)
        y += 34
        self.text(s, M, y, cw, 44, d["proposal"], 28, INK, bold=True)
        y += 46
        if d.get("proposal_text"):
            self.text(s, M, y, cw, 60, d["proposal_text"], 18, FG2, lh=1.38)
            y += 18 * 1.38 * min(2, _lines(d["proposal_text"], cw, 18)) + 22
        # how
        self.text(s, M, y, cw, 26, d["how_label"], 18, TEAL, bold=True, fit_box=False)
        y += 36
        how = d["how"][:5]
        hw = (cw - 12 * (len(how) - 1)) / len(how)
        for i, h in enumerate(how):
            x = M + i * (hw + 12)
            self.box(s, x, y, hw, 170, radius=14)
            self.text(s, x + 16, y + 18, hw - 32, 22, h["label"], 15, EDGE, bold=True, fit_box=False)
            self.text(s, x + 16, y + 48, hw - 32, 110, h["value"], 17, INK, bold=True, lh=1.28)
        y += 170 + 30
        # price band
        self.box(s, M, y, cw, 96, fill=TEAL_SOFT, line=None, radius=16)
        self.text(s, M + 30, y, cw * 0.55, 96, d["price"], 34, TEAL, bold=True, anchor="m", fit_box=False)
        self.text(s, M + cw * 0.55, y, cw * 0.45 - 30, 96, d.get("price_note", ""), 18, FG, anchor="m", align="r")
        y += 96 + 28
        self.text(s, M, y, cw, 26, d["next_label"], 18, TEAL, bold=True, fit_box=False)
        y += 34
        self.text(s, M, y, cw, 90, d["next"], 20, FG, lh=1.4)
        self.rule(s, M, PH - 92, cw)
        self.text(s, M, PH - 74, cw * 0.7, 24, d.get("contact", ""), 14, FG2, fit_box=False)
        self.text(s, M + cw * 0.6, PH - 74, cw * 0.4, 24, d.get("confidential", ""), 14, EDGE, align="r", fit_box=False)
        return s

    # ---------------------------------------------------------------- output
    def save(self, path):
        self.finish_notes()
        self.prs.save(path)
        return path


SLIDES = {
    "cover": Deck.cover, "statement": Deck.statement, "questions": Deck.questions,
    "what_we_do": Deck.what_we_do, "where_we_stand": Deck.where_we_stand,
    "coverage": Deck.coverage, "portfolio": Deck.portfolio, "outputs_impact": Deck.outputs_impact,
    "cooperations": Deck.cooperations, "engine": Deck.engine, "steps": Deck.steps,
    "catalog": Deck.catalog, "detail": Deck.detail_slide, "factor_pricing": Deck.factor_pricing,
    "build_vs_buy": Deck.build_vs_buy, "year_plan": Deck.year_plan, "understood": Deck.understood,
    "formats": Deck.formats, "blocks": Deck.blocks, "investment": Deck.investment,
    "sample": Deck.sample, "closing": Deck.closing, "onepager": Deck.onepager,
}
